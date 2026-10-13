import asyncio
import io
import os
import re
from datetime import datetime

import aiohttp
import pandas as pd
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from brands.vans import parse_vans_catalog, ALL_UNIQUE_SUFFIXES as VANS_SUFFIXES
from brands.timberland import (
    scrape_timberland_catalog,
    ALL_UNIQUE_SUFFIXES as TIMBERLAND_SUFFIXES,
    PROBE_SUFFIX as TIMBERLAND_PROBE,
)

app = FastAPI(title="Catalog Asset Intelligence API", docs_url=None, redoc_url=None, openapi_url=None)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
    expose_headers=["Content-Disposition"],
)

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
}

# Everything customers should not need to see or tune lives here, server-side.
BRAND_CONFIG = {
    "vans": {"base": "https://img.vans.com/image/upload/", "preset": ""},
    "timberland": {"base": "https://assets.timberland.com/", "preset": ""},
}
MAX_CONCURRENT = int(os.getenv("MAX_CONCURRENT", "400"))
REQUEST_TIMEOUT = int(os.getenv("REQUEST_TIMEOUT", "5"))

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ------------------------------------------------------------------ validation
async def check_url(session, semaphore, url, timeout):
    async with semaphore:
        t = aiohttp.ClientTimeout(total=timeout)
        for attempt in range(3):
            try:
                async with session.head(url, timeout=t, allow_redirects=True) as resp:
                    if resp.status == 200 and "x-cld-error" not in resp.headers:
                        ctype = resp.headers.get("Content-Type", "image").lower()
                        return url, ctype.startswith("image")
                    if resp.status in (404, 410):
                        return url, False
            except Exception:
                pass
            await asyncio.sleep(0.5 * (attempt + 1))
        return url, False


async def validate_urls(url_list, max_concurrent, timeout):
    semaphore = asyncio.Semaphore(max_concurrent)
    connector = aiohttp.TCPConnector(limit=max_concurrent, ttl_dns_cache=600, ssl=False)
    results = {}
    async with aiohttp.ClientSession(connector=connector, headers=HEADERS) as session:
        tasks = [check_url(session, semaphore, url, timeout) for url in url_list]
        for future in asyncio.as_completed(tasks):
            url, status = await future
            results[url] = status
    return results


# ------------------------------------------------------------------ excel helpers
HEADER_FILL = PatternFill("solid", fgColor="0F172A")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name="Calibri", size=10)
LINK_FONT = Font(name="Calibri", size=10, color="1D4ED8", underline="single")
THIN = Side(style="thin", color="E2E8F0")
BORDER = Border(bottom=THIN)
ACTIVE_FILL = PatternFill("solid", fgColor="DCFCE7")
MISSING_FILL = PatternFill("solid", fgColor="FEE2E2")


def style_sheet(ws, headers, link_cols=(), widths=None):
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(14, len(h) + 4))


def finish_sheet(ws, link_cols=()):
    last_col = get_column_letter(ws.max_column)
    ws.auto_filter.ref = f"A1:{last_col}{ws.max_row}"
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY_FONT
            c.border = BORDER
            c.alignment = Alignment(vertical="center")
    for col in link_cols:
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            for c in row:
                if c.value:
                    c.hyperlink = str(c.value)
                    c.font = LINK_FONT


def workbook_response(wb, filename):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type=XLSX_MIME,
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Cache-Control": "no-store",
        },
    )


def safe_filename(name, default):
    name = re.sub(r"[^A-Za-z0-9._-]+", "_", str(name or "")).strip("._")
    if not name:
        name = default
    return name if name.lower().endswith(".xlsx") else name + ".xlsx"


# ------------------------------------------------------------------ endpoints
@app.post("/api/v1/analyze")
async def analyze_catalog(brand: str = Form("vans"), file: UploadFile = File(...)):
    contents = await file.read()
    file_name = (file.filename or "").lower()

    try:
        if file_name.endswith(".csv"):
            df = pd.read_csv(io.BytesIO(contents), low_memory=False)
        elif file_name.endswith((".xlsx", ".xls")):
            df = pd.read_excel(io.BytesIO(contents))
        else:
            raise HTTPException(status_code=400, detail="Unsupported file format. Please upload a .csv or .xlsx file.")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=400, detail="The file could not be read. Please check that it is a valid spreadsheet.")

    df.columns = [str(c).strip() for c in df.columns]
    if "VENDOR ARTICLE NO." not in df.columns:
        raise HTTPException(status_code=400, detail="Required column 'VENDOR ARTICLE NO.' was not found in the file.")

    brand_lower = brand.lower().strip()
    cfg = BRAND_CONFIG.get(brand_lower)
    if not cfg:
        raise HTTPException(status_code=400, detail="Unsupported brand selection.")

    if brand_lower == "vans":
        for col in ("MH3 DIVISION", "AH1 SILLOUTE"):
            if col not in df.columns:
                df[col] = ""
        item_url_map, all_urls, _ = parse_vans_catalog(df, cfg["base"], cfg["preset"])
    else:
        item_url_map, all_urls, _ = scrape_timberland_catalog(df, cfg["base"], cfg["preset"])

    if not item_url_map:
        raise HTTPException(status_code=400, detail="No valid products were found in the file.")

    probe = TIMBERLAND_PROBE if brand_lower == "timberland" else None
    if probe:
        probe_urls = {d["urls"][probe] for d in item_url_map.values() if probe in d["urls"]}
        valid_url_map = await validate_urls(list(probe_urls), MAX_CONCURRENT, REQUEST_TIMEOUT)

        rest = set()
        for d in item_url_map.values():
            if valid_url_map.get(d["urls"].get(probe), False):
                rest.update(u for sfx, u in d["urls"].items() if sfx != probe)
        rest -= set(valid_url_map)
        valid_url_map.update(await validate_urls(list(rest), MAX_CONCURRENT, REQUEST_TIMEOUT))

        for u in all_urls:
            valid_url_map.setdefault(u, False)
    else:
        valid_url_map = await validate_urls(list(all_urls), MAX_CONCURRENT, REQUEST_TIMEOUT)

    return {
        "brand": brand_lower,
        "items": item_url_map,
        "available": valid_url_map,
        "summary": {
            "total_items": len(item_url_map),
            "active_items": sum(
                1 for v in item_url_map.values() if any(valid_url_map.get(u, False) for u in v["urls"].values())
            ),
            "total_images_found": sum(1 for ok in valid_url_map.values() if ok),
        },
    }


@app.post("/api/v1/reports/coverage")
async def coverage_report(payload: dict):
    brand = str(payload.get("brand", "vans")).lower()
    items = payload.get("items", {})
    available = payload.get("available", {})
    suffixes = TIMBERLAND_SUFFIXES if brand == "timberland" else VANS_SUFFIXES

    def title(sfx):
        return "Hero View" if sfx == "-HERO" else f"Alt {sfx.replace('-ALT', '')}"

    wb = Workbook()
    ws = wb.active
    ws.title = "Coverage Report"
    headers = ["Item Code", "Root Code", "Division", "Status", "Available Images"] + [title(s) for s in suffixes]
    widths = {"Item Code": 18, "Root Code": 18, "Division": 14, "Status": 12, "Available Images": 17}
    style_sheet(ws, headers, widths=widths)

    active = 0
    for code, data in items.items():
        urls = data.get("urls", {})
        valid = sum(1 for u in urls.values() if available.get(u, False))
        active += 1 if valid else 0
        row = [code, data.get("root", ""), data.get("division", ""), "Active" if valid else "Missing", valid]
        for s in suffixes:
            u = urls.get(s)
            row.append(u if u and available.get(u, False) else "")
        ws.append(row)

    link_cols = range(6, len(headers) + 1)
    finish_sheet(ws, link_cols=link_cols)
    for r in range(2, ws.max_row + 1):
        c = ws.cell(row=r, column=4)
        c.fill = ACTIVE_FILL if c.value == "Active" else MISSING_FILL
        c.font = Font(name="Calibri", size=10, bold=True, color="166534" if c.value == "Active" else "991B1B")
    for i in range(6, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 22

    total = len(items)
    summ = wb.create_sheet("Summary", 0)
    summ.sheet_view.showGridLines = False
    summ["B2"] = "Catalog Coverage Summary"
    summ["B2"].font = Font(name="Calibri", size=16, bold=True, color="0F172A")
    summ["B3"] = f"Generated {datetime.now().strftime('%d %b %Y, %H:%M')}"
    summ["B3"].font = Font(name="Calibri", size=10, color="64748B")
    rows = [
        ("Total products", total),
        ("Products with imagery", active),
        ("Products without imagery", total - active),
        ("Coverage", (active / total) if total else 0),
        ("Verified images", sum(1 for v in available.values() if v)),
    ]
    for i, (k, v) in enumerate(rows, start=5):
        summ.cell(row=i, column=2, value=k).font = Font(name="Calibri", size=11, color="334155")
        c = summ.cell(row=i, column=3, value=v)
        c.font = Font(name="Calibri", size=11, bold=True, color="0F172A")
        c.alignment = Alignment(horizontal="right")
        if k == "Coverage":
            c.number_format = "0.0%"
        for col in (2, 3):
            summ.cell(row=i, column=col).border = BORDER
    summ.column_dimensions["A"].width = 3
    summ.column_dimensions["B"].width = 30
    summ.column_dimensions["C"].width = 16

    return workbook_response(wb, "Catalog_Coverage_Report.xlsx")


@app.post("/api/v1/reports/export")
async def export_report(payload: dict):
    rows = payload.get("rows", [])
    if not isinstance(rows, list) or not rows:
        raise HTTPException(status_code=400, detail="There are no images to export.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Image Export"
    headers = ["Handle", "Item Code", "Image Position", "Image Src", "Image Alt Text", "Division"]
    style_sheet(ws, headers, widths={"Handle": 20, "Item Code": 18, "Image Position": 16, "Image Src": 70, "Image Alt Text": 32, "Division": 14})
    for r in rows:
        ws.append([
            r.get("Handle", ""),
            r.get("Item Code", ""),
            r.get("Image Position", ""),
            r.get("Image Src", ""),
            r.get("Image Alt Text", ""),
            r.get("Division", ""),
        ])
    finish_sheet(ws, link_cols=(4,))
    return workbook_response(wb, safe_filename(payload.get("filename"), "Image_Export"))


static_dir = os.path.join(os.path.dirname(__file__), "static")
if os.path.exists(static_dir):
    app.mount("/", StaticFiles(directory=static_dir, html=True), name="static")
